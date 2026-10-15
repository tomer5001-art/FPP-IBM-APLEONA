"""Builds Attendance_2026.xlsx - attendance, vacation and sick leave tracker for 2026.

Run: python3 excel/build_attendance_xlsx.py [output_path]
"""
import sys
from datetime import date, time, timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else 'Attendance_2026.xlsx'

EMPLOYEES = [('Tomer Cohen', 'Administrator'), ('Nir Ben Haim', 'Employee'), ('Jonatan Ben Sudai', 'Employee')]
STATUSES = ['Present', 'Vacation', 'Sick Leave', 'Partial Vacation', 'Partial Sick Leave', 'Manual Work', 'No Work']
MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
HOLIDAYS = [
    ('2026-04-01', 'Erev Pesach', 'Eve'), ('2026-04-02', 'Pesach', 'Holiday'),
    ('2026-04-07', 'Erev Seventh Day of Pesach', 'Eve'), ('2026-04-08', 'Seventh Day of Pesach', 'Holiday'),
    ('2026-04-22', 'Independence Day', 'Holiday'),
    ('2026-05-21', 'Erev Shavuot', 'Eve'), ('2026-05-22', 'Shavuot', 'Holiday'),
    ('2026-09-11', 'Erev Rosh Hashana', 'Eve'), ('2026-09-12', 'Rosh Hashana', 'Holiday'),
    ('2026-09-13', 'Rosh Hashana (second day)', 'Holiday'),
    ('2026-09-20', 'Erev Yom Kippur', 'Eve'), ('2026-09-21', 'Yom Kippur', 'Holiday'),
    ('2026-09-25', 'Erev Sukkot', 'Eve'), ('2026-09-26', 'Sukkot', 'Holiday'),
    ('2026-10-02', 'Erev Simchat Torah', 'Eve'), ('2026-10-03', 'Simchat Torah', 'Holiday'),
]

# Layout constants -----------------------------------------------------------
FIRST, LAST = 7, 371            # data rows on employee sheets (365 days)
HOL_FIRST, HOL_LAST = 5, 60     # holiday table rows (room to add more)
EMP_FIRST, EMP_LAST = 5, 14     # employee list rows (up to 10 employees)
S_START, S_END, S_EVE, S_NORMAL, S_EVEDAY, S_EVEVAC = (f'Settings!$B${r}' for r in (6, 7, 8, 9, 10, 11))
HOL_DATES = f'Holidays!$A${HOL_FIRST}:$A${HOL_LAST}'
HOL_NAMES = f'Holidays!$C${HOL_FIRST}:$C${HOL_LAST}'
HOL_TYPES = f'Holidays!$D${HOL_FIRST}:$D${HOL_LAST}'

# Styles -----------------------------------------------------------------------
NAVY = '1F3864'
FONT = Font(name='Arial', size=10)
BOLD = Font(name='Arial', size=10, bold=True)
TITLE = Font(name='Arial', size=14, bold=True, color=NAVY)
SUBTITLE = Font(name='Arial', size=10, italic=True, color='595959')
HEAD = Font(name='Arial', size=10, bold=True, color='FFFFFF')
INPUT_FONT = Font(name='Arial', size=10, color='0000FF')
LINK_FONT = Font(name='Arial', size=10, color='008000')
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
SUBHEAD_FILL = PatternFill('solid', fgColor='D9E1F2')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
TOTAL_FILL = PatternFill('solid', fgColor='F2F2F2')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_RULE = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')
RIGHT = Alignment(horizontal='right', vertical='center')
WRAP = Alignment(horizontal='left', vertical='center', wrap_text=True)
UNLOCKED = Protection(locked=False)

FMT_DATE = 'dd/mm/yyyy'
FMT_TIME = 'hh:mm'
FMT_NUM = '#,##0.00;(#,##0.00);"-"'
FMT_INT = '#,##0;(#,##0);"-"'

wb = Workbook()


def sheet(title):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    return ws


def style(c, font=FONT, fill=None, fmt=None, align=None, border=BOX):
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def header(ws, row, labels, col=1):
    for i, label in enumerate(labels):
        style(ws.cell(row, col + i, label), HEAD, HEAD_FILL, align=Alignment(horizontal='center', vertical='center', wrap_text=True))
    ws.row_dimensions[row].height = 42


def as_input(c, fmt=None, align=None):
    style(c, INPUT_FONT, INPUT_FILL, fmt, align)
    c.protection = UNLOCKED


def widths(ws, values):
    for i, w in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def protect(ws):
    p = ws.protection
    p.sheet = True
    p.formatColumns = False
    p.formatRows = False
    p.autoFilter = False
    p.sort = False


def title(ws, text, sub):
    ws['A1'] = text
    ws['A1'].font = TITLE
    ws['A2'] = sub
    ws['A2'].font = SUBTITLE


# Lists (hidden) -----------------------------------------------------------------
wl = wb.active
wl.title = 'Lists'
for i, s in enumerate(STATUSES, 1):
    wl.cell(i, 1, s)
for i, m in enumerate(MONTHS, 1):
    wl.cell(i, 2, m)
wl.cell(1, 3, 'Holiday'); wl.cell(2, 3, 'Eve')
wl.cell(1, 4, 'Administrator'); wl.cell(2, 4, 'Employee')
wl.cell(1, 5, 'Yes'); wl.cell(2, 5, 'No')
wl.sheet_state = 'hidden'

# Instructions -----------------------------------------------------------------
wi = sheet('Instructions')
title(wi, 'Attendance & Vacation Tracker 2026', 'Internal attendance, vacation and sick leave record. Calendar year 2026.')
lines = [
    ('How it works', None),
    ('Sunday - Thursday', 'Automatically Present from the default start to the default end time (08:00 - 16:30 = 8.5 hours). Nothing needs to be entered.'),
    ('Friday and Saturday', 'No automatic hours. Enter Manual Work with start and end time if work was done. Saturday work is entered by the administrator only.'),
    ('Holidays', 'No automatic work. Vacation and sick leave cannot be recorded on a holiday. Holidays are not counted as working days.'),
    ('Holiday eve', 'Automatic work until 14:00. A full vacation day on a holiday eve counts as half a vacation day.'),
    ('Partial leave', 'Day equivalent = leave hours / 8.5. Example: 4.25 hours = 0.50 day.'),
    ('', None),
    ('What to fill in', None),
    ('Yellow cells, blue text', 'Input cells. Everything else is calculated and protected.'),
    ('Status', 'Leave empty for an automatic day. Choose from the list only when the day is different.'),
    ('Start / End', 'Working hours (hh:mm). Required for Manual Work. Optional for partial leave (the hours actually worked).'),
    ('Leave from / Leave to', 'Required for Partial Vacation and Partial Sick Leave.'),
    ('Notes', 'Optional, e.g. Doctor appointment, Annual vacation, Personal leave.'),
    ('Check column', 'Shows a message when an entry is incomplete or not allowed. Such entries are not counted.'),
    ('', None),
    ('Adding an employee', None),
    ('Step 1', 'Right-click an employee sheet tab > Move or Copy > Create a copy.'),
    ('Step 2', 'Rename the new sheet to the employee full name, then clear the yellow input cells (columns E - J).'),
    ('Step 3', 'Add the same name to the Employees sheet. The reports and the dashboard pick it up automatically.'),
    ('', None),
    ('Sharing', None),
    ('Recommended', 'Keep the file on OneDrive or SharePoint so everyone edits the same copy. Version history there serves as the change log.'),
    ('Limitation', 'Excel cannot restrict access per employee: everyone with access to the file can see and edit every sheet.'),
]
r = 4
for label, text in lines:
    if text is None and label:
        style(wi.cell(r, 1, label), BOLD, SUBHEAD_FILL, border=None)
        wi.cell(r, 2).fill = SUBHEAD_FILL
    elif label:
        wi.cell(r, 1, label).font = BOLD
        wi.cell(r, 1).alignment = Alignment(vertical='center')
        wi.cell(r, 2, text).font = FONT
        wi.cell(r, 2).alignment = WRAP
    r += 1

r += 1
style(wi.cell(r, 1, 'Example entry'), BOLD, SUBHEAD_FILL, border=None)
wi.cell(r, 2).fill = SUBHEAD_FILL
r += 1
ex_head = ['Date', 'Status', 'Start', 'End', 'Leave from', 'Leave to', 'Notes', 'Worked hours', 'Vacation hours', 'Vacation days']
for i, h in enumerate(ex_head):
    style(wi.cell(r, 3 + i, h), HEAD, HEAD_FILL, align=CENTER)
wi.cell(r, 1, 'Partial vacation').font = BOLD
r += 1
wi.cell(r, 1, 'Works 08:00-12:00, vacation 12:00-16:30').font = FONT
ex = [date(2026, 1, 5), 'Partial Vacation', time(8, 0), time(12, 0), time(12, 0), time(16, 30), 'Doctor appointment']
fmts = [FMT_DATE, None, FMT_TIME, FMT_TIME, FMT_TIME, FMT_TIME, None]
for i, (v, f) in enumerate(zip(ex, fmts)):
    style(wi.cell(r, 3 + i, v), INPUT_FONT, INPUT_FILL, f, CENTER)
ex_row = r
style(wi.cell(r, 10, f'=(F{ex_row}-E{ex_row})*24'), fmt=FMT_NUM, align=RIGHT)
style(wi.cell(r, 11, f'=(H{ex_row}-G{ex_row})*24'), fmt=FMT_NUM, align=RIGHT)
style(wi.cell(r, 12, f'=K{ex_row}/{S_NORMAL}'), fmt=FMT_NUM, align=RIGHT)
wi.cell(r + 1, 1, 'Illustration only - the example is not part of any employee record.').font = SUBTITLE
widths(wi, [26, 16, 12, 18, 9, 9, 11, 11, 20, 14, 15, 14])
wi.column_dimensions['B'].width = 16
# instruction text spans B:L
for rr in range(4, ex_row - 2):
    if wi.cell(rr, 2).value and wi.cell(rr, 1).value and wi.cell(rr, 1).font != BOLD or True:
        pass
for rr in range(4, 4 + len(lines)):
    if wi.cell(rr, 2).value:
        wi.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=12)
        wi.row_dimensions[rr].height = 20
    elif wi.cell(rr, 1).value:
        wi.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=12)
protect(wi)

# Settings -----------------------------------------------------------------------
wst = sheet('Settings')
title(wst, 'Settings', 'Administrator only. Changes apply to the whole year.')
header(wst, 4, ['Parameter', 'Value', 'Note'])
settings_rows = [
    (5, 'Year', 2026, '0', 'Fixed.', False),
    (6, 'Default start time', time(8, 0), FMT_TIME, 'Applied to every automatic working day.', True),
    (7, 'Default end time', time(16, 30), FMT_TIME, 'Applied to every automatic working day.', True),
    (8, 'Holiday eve end time', time(14, 0), FMT_TIME, 'Working day ends at this time on a holiday eve.', True),
    (9, 'Normal working day (hours)', '=(B7-B6)*24', FMT_NUM, 'Calculated. Basis for partial leave day equivalents.', False),
    (10, 'Holiday eve working day (hours)', '=(B8-B6)*24', FMT_NUM, 'Calculated.', False),
    (11, 'Vacation day on a holiday eve (days)', 0.5, FMT_NUM, 'Company rule: a vacation day on a holiday eve counts as half a day.', True),
]
for row, label, val, fmt, note, editable in settings_rows:
    style(wst.cell(row, 1, label), BOLD)
    c = wst.cell(row, 2, val)
    if editable:
        as_input(c, fmt, RIGHT)
    else:
        style(c, fmt=fmt, align=RIGHT)
    style(wst.cell(row, 3, note), SUBTITLE)
widths(wst, [36, 12, 64])
protect(wst)

# Holidays -----------------------------------------------------------------------
wh = sheet('Holidays')
title(wh, 'Holidays 2026', 'Holiday: no work, no leave. Eve: short working day. Add rows below the list as needed.')
header(wh, 4, ['Date', 'Day', 'Holiday', 'Type'])
dv_type = DataValidation(type='list', formula1='=Lists!$C$1:$C$2', allow_blank=True)
dv_hdate = DataValidation(type='date', operator='between', formula1='DATE(2026,1,1)', formula2='DATE(2026,12,31)', allow_blank=True,
                          error='Enter a date in 2026.', showErrorMessage=True)
wh.add_data_validation(dv_type)
wh.add_data_validation(dv_hdate)
for i in range(HOL_LAST - HOL_FIRST + 1):
    row = HOL_FIRST + i
    h = HOLIDAYS[i] if i < len(HOLIDAYS) else None
    as_input(wh.cell(row, 1, date.fromisoformat(h[0]) if h else None), FMT_DATE, CENTER)
    style(wh.cell(row, 2, f'=IF(A{row}="","",CHOOSE(WEEKDAY(A{row}),"Sun","Mon","Tue","Wed","Thu","Fri","Sat"))'), align=CENTER)
    as_input(wh.cell(row, 3, h[1] if h else None), align=LEFT)
    as_input(wh.cell(row, 4, h[2] if h else None), align=CENTER)
dv_type.add(f'D{HOL_FIRST}:D{HOL_LAST}')
dv_hdate.add(f'A{HOL_FIRST}:A{HOL_LAST}')
wh.freeze_panes = 'A5'
widths(wh, [13, 8, 34, 12])
protect(wh)

# Employees ----------------------------------------------------------------------
we = sheet('Employees')
title(we, 'Employees', 'The name must match the employee sheet tab exactly.')
header(we, 4, ['Employee', 'Role', 'Active'])
dv_role = DataValidation(type='list', formula1='=Lists!$D$1:$D$2', allow_blank=True)
dv_yes = DataValidation(type='list', formula1='=Lists!$E$1:$E$2', allow_blank=True)
we.add_data_validation(dv_role)
we.add_data_validation(dv_yes)
for i in range(EMP_LAST - EMP_FIRST + 1):
    row = EMP_FIRST + i
    e = EMPLOYEES[i] if i < len(EMPLOYEES) else None
    as_input(we.cell(row, 1, e[0] if e else None), align=LEFT)
    as_input(we.cell(row, 2, e[1] if e else None), align=CENTER)
    as_input(we.cell(row, 3, 'Yes' if e else None), align=CENTER)
dv_role.add(f'B{EMP_FIRST}:B{EMP_LAST}')
dv_yes.add(f'C{EMP_FIRST}:C{EMP_LAST}')
widths(we, [28, 16, 10])
protect(we)

# Employee sheets ----------------------------------------------------------------
EMP_HEAD = ['Date', 'Day', 'Day type', 'Holiday', 'Status\n(leave empty = automatic)', 'Start', 'End', 'Leave from', 'Leave to', 'Notes',
            'Applied status', 'Applied start', 'Applied end', 'Worked hours', 'Vacation hours', 'Sick leave hours', 'Vacation days',
            'Sick leave days', 'Check', 'Month', 'Working day', 'Day end', 'Full day hours']


def employee_sheet(name):
    ws = sheet(name)
    title(ws, f'Attendance 2026 - {name}', 'Fill in only the yellow cells, and only for days that differ from the automatic day.')
    # summary block
    summary = [
        ('Vacation days', f'=SUM(Q{FIRST}:Q{LAST})', FMT_NUM),
        ('Vacation hours', f'=SUM(O{FIRST}:O{LAST})', FMT_NUM),
        ('Sick leave days', f'=SUM(R{FIRST}:R{LAST})', FMT_NUM),
        ('Sick leave hours', f'=SUM(P{FIRST}:P{LAST})', FMT_NUM),
        ('Days worked', f'=COUNTIF(N{FIRST}:N{LAST},">0")', FMT_INT),
        ('Worked hours', f'=SUM(N{FIRST}:N{LAST})', FMT_NUM),
        ('Entries to check', f'=COUNTIF(S{FIRST}:S{LAST},"?*")', FMT_INT),
    ]
    spans = [(1, 2), (3, 4), (5, 5), (6, 7), (8, 9), (10, 10), (11, 11)]
    for (label, f, fmt), (c1, c2) in zip(summary, spans):
        style(ws.cell(3, c1, label), BOLD, SUBHEAD_FILL, align=CENTER)
        style(ws.cell(4, c1, f), BOLD, fmt=fmt, align=CENTER)
        for c in range(c1 + 1, c2 + 1):
            style(ws.cell(3, c), BOLD, SUBHEAD_FILL)
            style(ws.cell(4, c))
        if c2 > c1:
            ws.merge_cells(start_row=3, start_column=c1, end_row=3, end_column=c2)
            ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)

    header(ws, 6, EMP_HEAD)
    for c in range(5, 11):
        ws.cell(6, c).fill = PatternFill('solid', fgColor='7F6000')
    for c in range(20, 24):
        ws.cell(6, c).fill = PatternFill('solid', fgColor='808080')

    dv_status = DataValidation(type='list', formula1='=Lists!$A$1:$A$7', allow_blank=True,
                               error='Choose a status from the list, or leave empty.', showErrorMessage=True)
    dv_time = DataValidation(type='time', operator='between', formula1='0', formula2='0.999305555', allow_blank=True,
                             error='Enter a time as hh:mm, e.g. 08:00.', showErrorMessage=True)
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_time)

    d0 = date(2026, 1, 1)
    for i in range(365):
        r = FIRST + i
        style(ws.cell(r, 1, d0 + timedelta(days=i)), fmt=FMT_DATE, align=CENTER)
        style(ws.cell(r, 2, f'=CHOOSE(WEEKDAY(A{r}),"Sun","Mon","Tue","Wed","Thu","Fri","Sat")'), align=CENTER)
        htype = f'IFERROR(INDEX({HOL_TYPES},MATCH(A{r},{HOL_DATES},0)),"")'
        style(ws.cell(r, 3, f'=IF({htype}="Holiday","Holiday",IF(WEEKDAY(A{r})>=6,"Weekend",IF({htype}="Eve","Eve","Workday")))'), align=CENTER)
        style(ws.cell(r, 4, f'=IFERROR(INDEX({HOL_NAMES},MATCH(A{r},{HOL_DATES},0)),"")'), LINK_FONT, align=LEFT)
        as_input(ws.cell(r, 5), align=LEFT)
        for c in (6, 7, 8, 9):
            as_input(ws.cell(r, c), FMT_TIME, CENTER)
        as_input(ws.cell(r, 10), align=LEFT)
        works = f'OR(K{r}="Present",K{r}="Manual Work",K{r}="Partial Vacation",K{r}="Partial Sick Leave")'
        formulas = {
            11: (f'=IF(E{r}<>"",E{r},IF(U{r},"Present",IF(C{r}="Holiday","Holiday","Weekend")))', None, LEFT),
            12: (f'=IF({works},IF(ISNUMBER(F{r}),F{r},IF(AND(K{r}="Present",U{r}),{S_START},"")),"")', FMT_TIME, CENTER),
            13: (f'=IF({works},IF(ISNUMBER(G{r}),G{r},IF(AND(K{r}="Present",U{r}),V{r},"")),"")', FMT_TIME, CENTER),
            14: (f'=IF(AND(ISNUMBER(L{r}),ISNUMBER(M{r})),MAX(0,(M{r}-L{r})*24),0)', FMT_NUM, RIGHT),
            15: (f'=IF(NOT(U{r}),0,IF(K{r}="Vacation",IF(C{r}="Eve",{S_NORMAL}*{S_EVEVAC},W{r}),'
                 f'IF(K{r}="Partial Vacation",IF(AND(ISNUMBER(H{r}),ISNUMBER(I{r})),MAX(0,(I{r}-H{r})*24),0),0)))', FMT_NUM, RIGHT),
            16: (f'=IF(NOT(U{r}),0,IF(K{r}="Sick Leave",W{r},'
                 f'IF(K{r}="Partial Sick Leave",IF(AND(ISNUMBER(H{r}),ISNUMBER(I{r})),MAX(0,(I{r}-H{r})*24),0),0)))', FMT_NUM, RIGHT),
            17: (f'=IF(NOT(U{r}),0,IF(K{r}="Vacation",IF(C{r}="Eve",{S_EVEVAC},1),IF(K{r}="Partial Vacation",O{r}/{S_NORMAL},0)))', FMT_NUM, RIGHT),
            18: (f'=IF(NOT(U{r}),0,IF(K{r}="Sick Leave",1,IF(K{r}="Partial Sick Leave",P{r}/{S_NORMAL},0)))', FMT_NUM, RIGHT),
            19: (f'=IF(AND(OR(E{r}="Vacation",E{r}="Sick Leave",E{r}="Partial Vacation",E{r}="Partial Sick Leave"),NOT(U{r})),"Leave not allowed on this day",'
                 f'IF(AND(E{r}="Manual Work",OR(NOT(ISNUMBER(F{r})),NOT(ISNUMBER(G{r})))),"Enter start and end",'
                 f'IF(AND(OR(E{r}="Partial Vacation",E{r}="Partial Sick Leave"),OR(NOT(ISNUMBER(H{r})),NOT(ISNUMBER(I{r})))),"Enter leave from and to",'
                 f'IF(AND(ISNUMBER(F{r}),ISNUMBER(G{r}),G{r}<=F{r}),"End is before start",'
                 f'IF(AND(ISNUMBER(H{r}),ISNUMBER(I{r}),I{r}<=H{r}),"Leave end is before start",'
                 f'IF(AND(ISNUMBER(F{r}),ISNUMBER(G{r}),ISNUMBER(H{r}),ISNUMBER(I{r}),F{r}<I{r},H{r}<G{r}),"Work and leave hours overlap",'
                 f'IF(AND(E{r}<>"",WEEKDAY(A{r})=7),"Saturday - administrator only","")))))))', Font(name='Arial', size=10, color='C00000'), LEFT),
            20: (f'=MONTH(A{r})', '0', CENTER),
            21: (f'=OR(C{r}="Workday",C{r}="Eve")', None, CENTER),
            22: (f'=IF(C{r}="Eve",{S_EVE},{S_END})', FMT_TIME, CENTER),
            23: (f'=IF(U{r},(V{r}-{S_START})*24,0)', FMT_NUM, RIGHT),
        }
        for c, (f, fmt, align) in formulas.items():
            if isinstance(fmt, Font):
                style(ws.cell(r, c, f), fmt, align=align)
            else:
                style(ws.cell(r, c, f), fmt=fmt, align=align)
    dv_status.add(f'E{FIRST}:E{LAST}')
    dv_time.add(f'F{FIRST}:I{LAST}')

    rng = f'A{FIRST}:S{LAST}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$C{FIRST}="Holiday"'], fill=PatternFill('solid', fgColor='E7E6E6'), font=Font(color='7F7F7F')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$C{FIRST}="Weekend"'], font=Font(color='A6A6A6')))
    ws.conditional_formatting.add(f'K{FIRST}:K{LAST}', FormulaRule(formula=[f'$E{FIRST}<>""'], font=Font(bold=True, color='1F3864')))

    widths(ws, [11, 6, 10, 22, 20, 8, 8, 10, 10, 24, 16, 9, 9, 9, 10, 10, 10, 10, 28, 7, 9, 8, 9])
    for col in ('T', 'U', 'V', 'W'):
        ws.column_dimensions[col].hidden = True
    ws.freeze_panes = 'E7'
    ws.auto_filter.ref = f'A6:S{LAST}'
    protect(ws)


for name, _ in EMPLOYEES:
    employee_sheet(name)


def emp_range(name_cell, col):
    return f'INDIRECT("\'"&{name_cell}&"\'!${col}${FIRST}:${col}${LAST}")'


# Dashboard ------------------------------------------------------------------------
wd = sheet('Dashboard')
title(wd, 'Dashboard 2026', 'Vacation and sick leave usage per employee. Updates automatically.')
wd['A3'] = 'Current month'
wd['A3'].font = BOLD
wd['J3'] = '=IF(YEAR(TODAY())=2026,MONTH(TODAY()),IF(YEAR(TODAY())<2026,1,12))'
wd['J3'].font = Font(name='Arial', size=10, color='FFFFFF')  # helper: month number
wd['B3'] = '=INDEX(Lists!$B$1:$B$12,J3)&" 2026"'
wd['B3'].font = BOLD
wd['D3'] = '=IF(YEAR(TODAY())=2026,TODAY(),DATE(2026,12,31))'
wd['D3'].number_format = '"As of "dd/mm/yyyy'
wd['D3'].font = SUBTITLE
wd.merge_cells('D3:F3')
header(wd, 5, ['Employee', 'Vacation days 2026', 'Vacation hours 2026', 'Sick leave days 2026', 'Sick leave hours 2026',
               'Days worked this month (to date)', 'Working days this month', 'Entries to check'])
for i in range(EMP_LAST - EMP_FIRST + 1):
    r = 6 + i
    n = f'$A{r}'
    style(wd.cell(r, 1, f'=IF(Employees!A{EMP_FIRST + i}="","",Employees!A{EMP_FIRST + i})'), LINK_FONT, align=LEFT)
    cells = [
        (2, f'SUM({emp_range(n, "Q")})', FMT_NUM),
        (3, f'SUM({emp_range(n, "O")})', FMT_NUM),
        (4, f'SUM({emp_range(n, "R")})', FMT_NUM),
        (5, f'SUM({emp_range(n, "P")})', FMT_NUM),
        (6, f'COUNTIFS({emp_range(n, "T")},$J$3,{emp_range(n, "N")},">0",{emp_range(n, "A")},"<="&$D$3)', FMT_INT),
        (7, f'COUNTIFS({emp_range(n, "T")},$J$3,{emp_range(n, "U")},TRUE)', FMT_INT),
        (8, f'COUNTIF({emp_range(n, "S")},"?*")', FMT_INT),
    ]
    for c, f, fmt in cells:
        style(wd.cell(r, c, f'=IF({n}="","",IFERROR({f},"Sheet missing"))'), fmt=fmt, align=RIGHT)
tot = 6 + EMP_LAST - EMP_FIRST + 1
style(wd.cell(tot, 1, 'Total'), BOLD, TOTAL_FILL, border=TOP_RULE)
for c in range(2, 9):
    col = get_column_letter(c)
    if c in (6, 7):
        style(wd.cell(tot, c), BOLD, TOTAL_FILL, border=TOP_RULE)
        continue
    style(wd.cell(tot, c, f'=SUM({col}6:{col}{tot - 1})'), BOLD, TOTAL_FILL, FMT_NUM if c < 6 else FMT_INT, RIGHT, TOP_RULE)

r0 = tot + 3
wd.cell(r0 - 1, 1, 'Vacation days by month').font = BOLD
header(wd, r0, ['Employee'] + [m[:3] for m in MONTHS] + ['Total'])
for i in range(EMP_LAST - EMP_FIRST + 1):
    r = r0 + 1 + i
    n = f'$A{r}'
    style(wd.cell(r, 1, f'=$A{6 + i}'), LINK_FONT, align=LEFT)
    for m in range(1, 13):
        style(wd.cell(r, 1 + m, f'=IF({n}="","",IFERROR(SUMIFS({emp_range(n, "Q")},{emp_range(n, "T")},{m}),"-"))'), fmt=FMT_NUM, align=RIGHT)
    style(wd.cell(r, 14, f'=IF({n}="","",SUM(B{r}:M{r}))'), BOLD, fmt=FMT_NUM, align=RIGHT)
widths(wd, [24, 12, 12, 12, 12, 14, 12, 10, 9, 9, 9, 9, 9, 10])
wd.freeze_panes = 'B6'
protect(wd)

# Monthly report -------------------------------------------------------------------
wm = sheet('Monthly Report')
title(wm, 'Monthly Attendance Report', 'Choose an employee and a month in the yellow cells.')
style(wm['A3'], BOLD, border=None); wm['A3'] = 'Employee'
style(wm['A4'], BOLD, border=None); wm['A4'] = 'Month'
as_input(wm['B3'], align=LEFT); wm['B3'] = EMPLOYEES[1][0]
as_input(wm['B4'], align=LEFT); wm['B4'] = 'January'
wm['D4'] = '=MATCH(B4,Lists!$B$1:$B$12,0)'
wm['D4'].font = Font(name='Arial', size=10, color='FFFFFF')  # helper: month number
wm.merge_cells('B3:C3'); wm.merge_cells('B4:C4')
dv_emp = DataValidation(type='list', formula1=f'=Employees!$A${EMP_FIRST}:$A${EMP_LAST}', allow_blank=False)
dv_month = DataValidation(type='list', formula1='=Lists!$B$1:$B$12', allow_blank=False)
wm.add_data_validation(dv_emp); wm.add_data_validation(dv_month)
dv_emp.add('B3'); dv_month.add('B4')
style(wm['E3'], BOLD, border=None)
wm['E3'] = '=B3&" - "&B4&" 2026"'

MR_HEAD = ['Date', 'Day', 'Status', 'Start', 'End', 'Total working hours', 'Vacation hours', 'Sick leave hours', 'Vacation day equivalent', 'Sick day equivalent', 'Notes']
header(wm, 6, MR_HEAD)
SRC = f'INDIRECT("\'"&$B$3&"\'!$A${FIRST}:$W${LAST}")'
for n in range(1, 32):
    r = 6 + n
    style(wm.cell(r, 1, f'=IF({n}<=DAY(DATE(2026,$D$4+1,0)),DATE(2026,$D$4,{n}),"")'), fmt=FMT_DATE, align=CENTER)
    pick = lambda col: f'INDEX({SRC},$A{r}-DATE(2026,1,1)+1,{col})'
    cols = [(2, pick(2), None, CENTER), (3, pick(11), None, LEFT), (4, pick(12), FMT_TIME, CENTER), (5, pick(13), FMT_TIME, CENTER),
            (6, pick(14), FMT_NUM, RIGHT), (7, pick(15), FMT_NUM, RIGHT), (8, pick(16), FMT_NUM, RIGHT),
            (9, pick(17), FMT_NUM, RIGHT), (10, pick(18), FMT_NUM, RIGHT), (11, f'""&{pick(10)}', None, LEFT)]
    for c, f, fmt, align in cols:
        style(wm.cell(r, c, f'=IF($A{r}="","",IFERROR({f},""))'), fmt=fmt, align=align)
wm.conditional_formatting.add('A7:K37', FormulaRule(formula=['OR($C7="Weekend",$C7="Holiday")'], font=Font(color='A6A6A6')))

t0 = 39
style(wm.cell(t0, 1, 'Monthly totals'), BOLD, SUBHEAD_FILL, border=None)
for c in range(2, 4):
    wm.cell(t0, c).fill = SUBHEAD_FILL
totals = [
    ('Total working days', '=COUNTIF(F7:F37,">0")', FMT_INT),
    ('Total worked hours', '=SUM(F7:F37)', FMT_NUM),
    ('Total vacation days', '=SUM(I7:I37)', FMT_NUM),
    ('Total sick leave days', '=SUM(J7:J37)', FMT_NUM),
    ('Total vacation hours', '=SUM(G7:G37)', FMT_NUM),
    ('Total sick leave hours', '=SUM(H7:H37)', FMT_NUM),
]
for i, (label, f, fmt) in enumerate(totals):
    r = t0 + 1 + i
    style(wm.cell(r, 1, label), BOLD)
    wm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    style(wm.cell(r, 3, f), BOLD, TOTAL_FILL, fmt, RIGHT)
widths(wm, [12, 8, 18, 8, 8, 12, 11, 11, 13, 12, 28])
wm.freeze_panes = 'A7'
wm.page_setup.orientation = 'landscape'
wm.page_setup.fitToWidth = 1
wm.page_setup.fitToHeight = 0
wm.sheet_properties.pageSetUpPr.fitToPage = True
wm.print_area = f'A1:K{t0 + len(totals)}'
protect(wm)

# Annual report --------------------------------------------------------------------
wa = sheet('Annual Report')
title(wa, 'Annual Report 2026', 'Per employee, 1 January - 31 December 2026.')
header(wa, 4, ['Employee', 'Total normal working days', 'Holidays (Sun-Thu)', 'Total days worked', 'Total worked hours',
               'Total vacation days taken', 'Total vacation hours', 'Total sick leave days', 'Total sick leave hours', 'Total manual work days'])
for i in range(EMP_LAST - EMP_FIRST + 1):
    r = 5 + i
    n = f'$A{r}'
    style(wa.cell(r, 1, f'=IF(Employees!A{EMP_FIRST + i}="","",Employees!A{EMP_FIRST + i})'), LINK_FONT, align=LEFT)
    cells = [
        (2, f'COUNTIF({emp_range(n, "U")},TRUE)', FMT_INT),
        (3, f'SUMPRODUCT(({emp_range(n, "C")}="Holiday")*(WEEKDAY({emp_range(n, "A")})<=5))', FMT_INT),
        (4, f'COUNTIF({emp_range(n, "N")},">0")', FMT_INT),
        (5, f'SUM({emp_range(n, "N")})', FMT_NUM),
        (6, f'SUM({emp_range(n, "Q")})', FMT_NUM),
        (7, f'SUM({emp_range(n, "O")})', FMT_NUM),
        (8, f'SUM({emp_range(n, "R")})', FMT_NUM),
        (9, f'SUM({emp_range(n, "P")})', FMT_NUM),
        (10, f'COUNTIF({emp_range(n, "K")},"Manual Work")', FMT_INT),
    ]
    for c, f, fmt in cells:
        style(wa.cell(r, c, f'=IF({n}="","",IFERROR({f},"Sheet missing"))'), fmt=fmt, align=RIGHT)
tot = 5 + EMP_LAST - EMP_FIRST + 1
style(wa.cell(tot, 1, 'Total'), BOLD, TOTAL_FILL, border=TOP_RULE)
for c in range(2, 11):
    col = get_column_letter(c)
    if c in (2, 3):
        style(wa.cell(tot, c), BOLD, TOTAL_FILL, border=TOP_RULE)
        continue
    style(wa.cell(tot, c, f'=SUM({col}5:{col}{tot - 1})'), BOLD, TOTAL_FILL, FMT_INT if c in (4, 10) else FMT_NUM, RIGHT, TOP_RULE)
widths(wa, [24, 13, 12, 12, 12, 13, 12, 12, 12, 12])
wa.freeze_panes = 'B5'
wa.page_setup.orientation = 'landscape'
wa.sheet_properties.pageSetUpPr.fitToPage = True
protect(wa)

# Order and defaults ---------------------------------------------------------------
order = ['Instructions', 'Dashboard'] + [e[0] for e in EMPLOYEES] + ['Monthly Report', 'Annual Report', 'Holidays', 'Employees', 'Settings', 'Lists']
wb._sheets = [wb[n] for n in order]
wb.active = 1
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = ws.title == 'Dashboard'
wb.save(OUT)
print('saved', OUT)
